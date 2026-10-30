from docx import Document
import json

doc = Document(r"C:\Users\User\Downloads\bangladesh_60_floor_plans_links.docx")

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(i, repr(p.text[:300]))

print("\n=== TABLES ===")
for ti, t in enumerate(doc.tables):
    print(f"TABLE {ti}")
    for ri, r in enumerate(t.rows):
        cells = []
        for c in r.cells:
            cells.append(c.text[:200])
        print(ri, cells)